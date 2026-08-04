import csv
import io
import re
import zipfile
from dataclasses import dataclass
from itertools import islice
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

MAX_QUESTIONS = 2000
MAX_COLUMNS = 250
MAX_SHEETS = 50
MAX_XLSX_ENTRIES = 2000
MAX_XLSX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ZIP_RATIO = 200
MAX_CELL_CHARS = 20_000
MAX_HEADER_CHARS = 255
QUESTION_HINTS = ("question", "requirement", "exigence", "contrôle", "control", "demande")


class QuestionnaireParseError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedSheet:
    source_format: str
    sheet_names: list[str]
    selected_sheet: str | None
    columns: list[str]
    rows: list[dict[str, Any]]
    suggested_question_column: str

    def questions(self, column: str | None = None) -> list[tuple[int, str, dict[str, Any]]]:
        selected = column or self.suggested_question_column
        if selected not in self.columns:
            raise QuestionnaireParseError("La colonne de questions sélectionnée n'existe pas")
        result: list[tuple[int, str, dict[str, Any]]] = []
        for row_number, row in enumerate(self.rows, start=1):
            raw = row.get(selected)
            if raw is None:
                continue
            text = str(raw).strip()
            if len(text) < 3:
                continue
            result.append((row_number, text[:20000], row))
            if len(result) >= MAX_QUESTIONS:
                break
        if not result:
            raise QuestionnaireParseError("Aucune question exploitable n'a été détectée")
        return result


def _unique_headers(values: list[Any]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(values, start=1):
        base = str(value).strip() if value is not None else ""
        if len(base) > MAX_HEADER_CHARS:
            raise QuestionnaireParseError(
                f"L'en-tête de la colonne {index} dépasse {MAX_HEADER_CHARS} caractères"
            )
        base = base or f"Colonne {index}"
        count = seen.get(base, 0) + 1
        seen[base] = count
        headers.append(base if count == 1 else f"{base} ({count})")
    return headers


def _question_score(header: str, values: list[Any]) -> float:
    normalized = header.casefold()
    score = 20.0 if any(hint in normalized for hint in QUESTION_HINTS) else 0.0
    populated = [str(value).strip() for value in values if value is not None and str(value).strip()]
    if not populated:
        return score
    question_marks = sum("?" in value for value in populated)
    long_text = sum(len(value) >= 25 for value in populated)
    return score + (question_marks / len(populated) * 10) + (long_text / len(populated) * 5)


def _choose_column(columns: list[str], rows: list[dict[str, Any]]) -> str:
    if not columns:
        raise QuestionnaireParseError("Le fichier ne contient aucune colonne")
    return max(
        columns, key=lambda column: _question_score(column, [row.get(column) for row in rows])
    )


def _json_value(value: Any) -> Any:
    if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
        raise QuestionnaireParseError(
            f"Une cellule dépasse la limite de {MAX_CELL_CHARS} caractères"
        )
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    rendered = str(value)
    if len(rendered) > MAX_CELL_CHARS:
        raise QuestionnaireParseError(
            f"Une cellule dépasse la limite de {MAX_CELL_CHARS} caractères"
        )
    return rendered


def _parse_csv(content: bytes) -> ParsedSheet:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise QuestionnaireParseError("Le CSV doit être encodé en UTF-8") from exc
    if "\x00" in text:
        raise QuestionnaireParseError("Le fichier CSV contient des octets invalides")
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    try:
        raw_rows = list(islice(reader, MAX_QUESTIONS + 2))
    except csv.Error as exc:
        raise QuestionnaireParseError(
            "Le fichier CSV contient un champ invalide ou trop long"
        ) from exc
    if len(raw_rows) > MAX_QUESTIONS + 1:
        raise QuestionnaireParseError(f"Le CSV dépasse la limite de {MAX_QUESTIONS} questions")
    if not raw_rows:
        raise QuestionnaireParseError("Le fichier CSV est vide")
    columns = _unique_headers(raw_rows[0])
    if len(columns) > MAX_COLUMNS:
        raise QuestionnaireParseError(f"Le CSV dépasse la limite de {MAX_COLUMNS} colonnes")
    rows = [
        {
            column: _json_value(values[index] if index < len(values) else None)
            for index, column in enumerate(columns)
        }
        for values in raw_rows[1:]
        if any(str(value).strip() for value in values)
    ]
    column = _choose_column(columns, rows)
    return ParsedSheet("csv", [], None, columns, rows, column)


def _sheet_score(rows: list[list[Any]]) -> float:
    if not rows:
        return -1
    headers = _unique_headers(rows[0])
    objects = [dict(zip(headers, row, strict=False)) for row in rows[1:100]]
    return max(
        (_question_score(header, [item.get(header) for item in objects]) for header in headers),
        default=-1,
    )


def _parse_xlsx(content: bytes, requested_sheet: str | None) -> ParsedSheet:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ENTRIES:
                raise QuestionnaireParseError("Le classeur contient trop d'entrées compressées")
            total_uncompressed = 0
            for entry in entries:
                if entry.flag_bits & 0x1:
                    raise QuestionnaireParseError("Les classeurs chiffrés ne sont pas acceptés")
                total_uncompressed += entry.file_size
                if total_uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise QuestionnaireParseError("Le classeur décompressé est trop volumineux")
                if entry.compress_size == 0 and entry.file_size > 0:
                    raise QuestionnaireParseError("Archive XLSX suspecte")
                if entry.compress_size and entry.file_size / entry.compress_size > MAX_ZIP_RATIO:
                    raise QuestionnaireParseError("Taux de compression XLSX suspect")
    except zipfile.BadZipFile as exc:
        raise QuestionnaireParseError("Le classeur XLSX est invalide") from exc
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise QuestionnaireParseError("Le classeur XLSX est invalide") from exc
    sheet_names = [name for name in workbook.sheetnames if workbook[name].sheet_state == "visible"]
    if not sheet_names:
        raise QuestionnaireParseError("Le classeur ne contient aucune feuille visible")
    if len(sheet_names) > MAX_SHEETS:
        raise QuestionnaireParseError(f"Le classeur dépasse la limite de {MAX_SHEETS} feuilles")
    if requested_sheet and requested_sheet not in sheet_names:
        raise QuestionnaireParseError("La feuille sélectionnée n'existe pas")
    cached: dict[str, list[list[Any]]] = {}
    for name in sheet_names:
        sheet = workbook[name]
        if sheet.max_column > MAX_COLUMNS:
            raise QuestionnaireParseError(
                f"La feuille « {name} » dépasse la limite de {MAX_COLUMNS} colonnes"
            )
        if sheet.max_row > MAX_QUESTIONS + 1:
            raise QuestionnaireParseError(
                f"La feuille « {name} » dépasse la limite de {MAX_QUESTIONS} questions"
            )
        cached[name] = [
            [_json_value(value) for value in row]
            for row in sheet.iter_rows(
                values_only=True,
                max_row=MAX_QUESTIONS + 1,
                max_col=max(1, sheet.max_column),
            )
        ]
    selected = requested_sheet or max(sheet_names, key=lambda name: _sheet_score(cached[name]))
    raw_rows = cached[selected]
    while raw_rows and not any(value is not None and str(value).strip() for value in raw_rows[0]):
        raw_rows.pop(0)
    if not raw_rows:
        raise QuestionnaireParseError("La feuille sélectionnée est vide")
    columns = _unique_headers(raw_rows[0])
    rows = [
        {
            column: _json_value(values[index] if index < len(values) else None)
            for index, column in enumerate(columns)
        }
        for values in raw_rows[1:]
        if any(value is not None and str(value).strip() for value in values)
    ]
    question_column = _choose_column(columns, rows)
    return ParsedSheet("xlsx", sheet_names, selected, columns, rows, question_column)


def parse_questionnaire(
    filename: str, content: bytes, sheet_name: str | None = None
) -> ParsedSheet:
    extension = Path(filename).suffix.casefold()
    if extension == ".csv":
        return _parse_csv(content)
    if extension == ".xlsx":
        return _parse_xlsx(content, sheet_name)
    raise QuestionnaireParseError("Seuls les formats CSV et XLSX sont pris en charge")


def looks_like_prompt_injection(value: str) -> bool:
    pattern = re.compile(
        r"(?i)(ignore\s+(all\s+)?(previous|prior|system)|system\s+prompt|developer\s+message|"
        r"révèle\s+.*secret|oublie\s+les\s+instructions)"
    )
    return bool(pattern.search(value))
