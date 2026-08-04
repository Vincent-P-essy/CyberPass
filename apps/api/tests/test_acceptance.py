import io
import json
import uuid
from datetime import UTC, datetime, timedelta

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.routes.auth import settings as auth_settings
from app.core.security import hash_token
from app.main import app as fastapi_app
from app.models import (
    AiUsageEvent,
    AuditEvent,
    Questionnaire,
    QuestionnaireState,
    ShareLink,
    SuggestedAnswer,
)
from app.services.ai import GenerationResult

PASSWORD = "Aaaaaaaaaaaa1"  # noqa: S105


def register(client: TestClient, email: str) -> tuple[str, str]:
    response = client.post(
        "/api/v1/auth/register",
        json={"email": email, "fullName": "Plessy Vincent", "password": PASSWORD},
    )
    assert response.status_code == 201, response.text
    return response.json()["accessToken"], response.json()["csrfToken"]


def auth(token: str, organization_id: str | None = None) -> dict[str, str]:
    headers = {"Authorization": f"Bearer {token}"}
    if organization_id:
        headers["X-Organization-ID"] = organization_id
    return headers


def create_organization(client: TestClient, token: str, name: str) -> str:
    response = client.post(
        "/api/v1/organizations",
        headers=auth(token),
        json={"name": name, "description": "Organisation fictive de test"},
    )
    assert response.status_code == 201, response.text
    return response.json()["id"]


def first_control(client: TestClient, token: str, organization_id: str) -> str:
    response = client.get("/api/v1/controls", headers=auth(token, organization_id))
    assert response.status_code == 200, response.text
    assert len(response.json()) == 15
    return response.json()[0]["id"]


def create_evidence(
    client: TestClient,
    token: str,
    organization_id: str,
    control_id: str,
    *,
    title: str = "Politique MFA",
    confidentiality: str = "PUBLIC",
    description: str = "Le MFA est activé pour les comptes administrateurs.",
    public_summary: str | None = None,
    with_file: bool = False,
) -> dict[str, object]:
    data = {
        "title": title,
        "evidenceType": "POLICY",
        "confidentiality": confidentiality,
        "description": description,
        "controlIds": json.dumps([control_id]),
    }
    if public_summary:
        data["publicSummary"] = public_summary
    files = (
        {"file": ("mfa-policy.txt", b"MFA enabled for administrators", "text/plain")}
        if with_file
        else None
    )
    response = client.post(
        "/api/v1/evidences",
        headers=auth(token, organization_id),
        data=data,
        files=files,
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_scenario_a_strict_tenant_isolation(client: TestClient) -> None:
    token_a, _csrf_a = register(client, "tenant-a-owner@cyberpass.example.com")
    organization_a = create_organization(client, token_a, "Organisation A")
    control_a = first_control(client, token_a, organization_a)
    evidence = create_evidence(client, token_a, organization_a, control_a)

    token_b, _csrf_b = register(client, "tenant-b-owner@cyberpass.example.com")
    organization_b = create_organization(client, token_b, "Organisation B")

    allowed = client.get(
        f"/api/v1/evidences/{evidence['id']}", headers=auth(token_a, organization_a)
    )
    assert allowed.status_code == 200

    denied = client.get(
        f"/api/v1/evidences/{evidence['id']}", headers=auth(token_b, organization_b)
    )
    assert denied.status_code == 404
    assert denied.json() == {"detail": "Preuve introuvable"}
    denied_mutation = client.patch(
        f"/api/v1/evidences/{evidence['id']}",
        headers=auth(token_b, organization_b),
        json={"title": "Tentative inter-tenant"},
    )
    assert denied_mutation.status_code == 404
    assert denied_mutation.json() == {"detail": "Preuve introuvable"}


def test_scenario_b_questionnaire_generation_review_and_export(client: TestClient) -> None:
    token, _csrf = register(client, "questionnaire-owner@cyberpass.example.com")
    organization_id = create_organization(client, token, "Questionnaire SAS")
    control_id = first_control(client, token, organization_id)
    evidence = create_evidence(client, token, organization_id, control_id)
    source_workbook = Workbook()
    source_sheet = source_workbook.active
    source_sheet.title = "Questionnaire sécurité"
    source_sheet.append(["Référence", "Question", "Commentaire"])
    source_sheet.append(
        ["Q1", "L'authentification multifacteur est-elle imposée aux administrateurs ?", ""]
    )
    source_sheet.append(["Q2", "Comment les accès sont-ils revus ?", ""])
    source_stream = io.BytesIO()
    source_workbook.save(source_stream)
    xlsx_content = source_stream.getvalue()

    preview = client.post(
        "/api/v1/questionnaires/preview",
        headers=auth(token, organization_id),
        files={
            "file": (
                "questionnaire.xlsx",
                xlsx_content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["suggestedQuestionColumn"] == "Question"
    assert preview.json()["detectedQuestionCount"] == 2

    imported = client.post(
        "/api/v1/questionnaires/import",
        headers=auth(token, organization_id),
        data={"name": "Questionnaire acheteur", "questionColumn": "Question"},
        files={
            "file": (
                "questionnaire.xlsx",
                xlsx_content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert imported.status_code == 201, imported.text
    questionnaire_id = imported.json()["id"]

    generated = client.post(
        f"/api/v1/questionnaires/{questionnaire_id}/generate",
        headers=auth(token, organization_id),
        json={"allowExternalProvider": False},
    )
    assert generated.status_code == 200, generated.text
    question = generated.json()["questions"][0]
    assert question["answerStatus"] == "GENERATED"
    assert question["answer"]["requiresHumanReview"] is True
    assert question["answer"]["evidenceIds"] == [evidence["id"]]

    edited = client.patch(
        f"/api/v1/questionnaires/{questionnaire_id}/questions/{question['id']}/answer",
        headers=auth(token, organization_id),
        json={"answer": "Oui. Le MFA est imposé aux administrateurs et revu trimestriellement."},
    )
    assert edited.status_code == 200, edited.text
    assert edited.json()["answerStatus"] == "MANUALLY_ANSWERED"

    approved = client.post(
        f"/api/v1/questionnaires/{questionnaire_id}/questions/{question['id']}/approve",
        headers=auth(token, organization_id),
        json={},
    )
    assert approved.status_code == 200, approved.text
    assert approved.json()["answerStatus"] == "APPROVED"
    assert approved.json()["answer"]["approvedAt"] is not None

    exported = client.get(
        f"/api/v1/questionnaires/{questionnaire_id}/export?format=xlsx",
        headers=auth(token, organization_id),
    )
    assert exported.status_code == 200, exported.text
    workbook = load_workbook(io.BytesIO(exported.content), read_only=True)
    exported_rows = list(workbook.active.iter_rows(values_only=True))
    assert exported_rows[0][:4] == ("Ordre", "Question", "Réponse", "Statut")
    assert "MFA est imposé" in exported_rows[1][2]
    after_read_only_export = client.get(
        f"/api/v1/questionnaires/{questionnaire_id}", headers=auth(token, organization_id)
    )
    assert after_read_only_export.json()["state"] == "IN_REVIEW"


def test_scenarios_c_and_d_share_allowlist_revocation_and_confidentiality(
    client: TestClient, db: Session
) -> None:
    token, _csrf = register(client, "passport-owner@cyberpass.example.com")
    organization_id = create_organization(client, token, "Passeport SARL")
    control_id = first_control(client, token, organization_id)
    public_evidence = create_evidence(
        client,
        token,
        organization_id,
        control_id,
        title="Synthèse MFA partageable",
        confidentiality="SHARED_SUMMARY",
        description="Détail interne à ne pas afficher",
        public_summary="Le MFA couvre les comptes d'administration.",
    )
    confidential = create_evidence(
        client,
        token,
        organization_id,
        control_id,
        title="SECRET-NOM-FICHIER",
        confidentiality="CONFIDENTIAL",
        description="SECRET-DETAIL-TECHNIQUE",
        with_file=True,
    )
    assert confidential["originalFilename"] == "mfa-policy.txt"
    verified = client.patch(
        f"/api/v1/controls/{control_id}",
        headers=auth(token, organization_id),
        json={"status": "VERIFIED"},
    )
    assert verified.status_code == 200, verified.text
    assert verified.json()["lastVerifiedAt"] is not None
    private_download = client.get(
        f"/api/v1/evidences/{confidential['id']}/download",
        headers=auth(token, organization_id),
    )
    assert private_download.status_code == 200
    assert private_download.headers["cache-control"] == "private, no-store"

    expires_at = (datetime.now(UTC) + timedelta(days=7)).isoformat()
    created = client.post(
        "/api/v1/share-links",
        headers=auth(token, organization_id),
        json={
            "title": "Passeport limité",
            "expiresAt": expires_at,
            "controlIds": [control_id],
            "evidenceIds": [public_evidence["id"]],
        },
    )
    assert created.status_code == 201, created.text
    share = created.json()
    stored = db.scalar(select(ShareLink).where(ShareLink.id == uuid.UUID(share["id"])))
    assert stored is not None
    assert stored.token_hash == hash_token(share["token"])
    assert share["token"] not in stored.token_hash

    public = client.get(share["publicPath"])
    assert public.status_code == 200, public.text
    payload = public.json()
    assert "title" not in payload
    assert [item["id"] for item in payload["controls"]] == [control_id]
    assert payload["controls"][0]["status"] == "VERIFIED"
    assert payload["controls"][0]["lastVerifiedAt"] is not None
    assert [item["title"] for item in payload["controls"][0]["evidences"]] == [
        "Résumé de preuve partagé"
    ]
    assert "id" not in payload["controls"][0]["evidences"][0]
    serialized = public.text
    assert "Synthèse MFA partageable" not in serialized
    assert "Le MFA couvre" in serialized
    assert "Détail interne" not in serialized
    assert "SECRET-NOM-FICHIER" not in serialized
    assert "SECRET-DETAIL-TECHNIQUE" not in serialized
    assert "mfa-policy.txt" not in serialized
    assert "objectKey" not in serialized
    assert "sha256" not in serialized
    assert "Passeport limité" not in serialized
    assert "certification" in payload["disclaimer"].lower()

    revoked = client.post(
        f"/api/v1/share-links/{share['id']}/revoke",
        headers=auth(token, organization_id),
    )
    assert revoked.status_code == 200, revoked.text
    unavailable = client.get(share["publicPath"])
    assert unavailable.status_code == 404


def test_audit_trail_and_csrf_double_submit(client: TestClient, db: Session) -> None:
    token, csrf_token = register(client, "csrf-owner@cyberpass.example.com")
    rejected = client.post("/api/v1/organizations", json={"name": "Sans CSRF"})
    assert rejected.status_code == 403

    accepted = client.post(
        "/api/v1/organizations",
        headers={"X-CSRF-Token": csrf_token},
        json={"name": "Avec CSRF"},
    )
    assert accepted.status_code == 201, accepted.text
    assert accepted.headers["cache-control"] == "private, no-store"
    assert accepted.headers["pragma"] == "no-cache"
    organization_id = accepted.json()["id"]
    events = client.get("/api/v1/audit-events", headers=auth(token, organization_id))
    assert events.status_code == 200
    assert "organization.created" in {event["action"] for event in events.json()}
    dashboard = client.get("/api/v1/dashboard", headers=auth(token, organization_id))
    assert dashboard.status_code == 200
    assert all(
        "userId" not in activity and "ipAddress" not in activity
        for activity in dashboard.json()["recentActivity"]
    )
    assert (
        db.scalar(
            select(AuditEvent).where(AuditEvent.organization_id == uuid.UUID(organization_id))
        )
        is not None
    )


def test_generation_batch_is_bounded(client: TestClient) -> None:
    token, _csrf = register(client, "batch-owner@cyberpass.example.com")
    organization_id = create_organization(client, token, "Batch limité")
    rows = ["Question"] + [f"La mesure numéro {index} est-elle appliquée ?" for index in range(26)]
    content = "\n".join(rows).encode()
    imported = client.post(
        "/api/v1/questionnaires/import",
        headers=auth(token, organization_id),
        files={"file": ("batch.csv", content, "text/csv")},
    )
    assert imported.status_code == 201, imported.text
    generated = client.post(
        f"/api/v1/questionnaires/{imported.json()['id']}/generate",
        headers=auth(token, organization_id),
        json={"allowExternalProvider": False},
    )
    assert generated.status_code == 422
    assert "25 questions" in generated.json()["detail"]


def test_generation_failure_is_atomic_and_audited(
    client: TestClient, db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    token, _csrf = register(client, "failure-owner@cyberpass.example.com")
    organization_id = create_organization(client, token, "Échec atomique")
    content = b"Question\nPremiere question ?\nDeuxieme question ?"
    imported = client.post(
        "/api/v1/questionnaires/import",
        headers=auth(token, organization_id),
        files={"file": ("failure.csv", content, "text/csv")},
    )
    assert imported.status_code == 201, imported.text

    class FailingProvider:
        provider_name = "test-provider"
        model_version = "test-model"

        def __init__(self) -> None:
            self.calls = 0

        def generate(self, question: str, controls: list[object], evidences: list[object]):
            self.calls += 1
            if self.calls == 2:
                raise RuntimeError("provider unavailable")
            return GenerationResult(
                proposed_answer=f"Réponse pour {question}",
                confidence=0.4,
                provider=self.provider_name,
                model_version=self.model_version,
            )

    provider = FailingProvider()
    monkeypatch.setattr(
        "app.api.routes.questionnaires.build_ai_provider", lambda _allow_external: provider
    )
    generated = client.post(
        f"/api/v1/questionnaires/{imported.json()['id']}/generate",
        headers=auth(token, organization_id),
        json={"allowExternalProvider": True},
    )
    assert generated.status_code == 502
    db.expire_all()
    questionnaire = db.get(Questionnaire, uuid.UUID(imported.json()["id"]))
    assert questionnaire is not None
    assert questionnaire.state == QuestionnaireState.FAILED
    assert db.scalars(select(SuggestedAnswer)).all() == []
    usages = db.scalars(select(AiUsageEvent).order_by(AiUsageEvent.created_at)).all()
    assert [usage.status for usage in usages] == ["SUCCESS", "FAILED"]
    assert usages[-1].error_code == "PROVIDER_ERROR"
    assert (
        db.scalar(select(AuditEvent).where(AuditEvent.action == "questionnaire.generation_failed"))
        is not None
    )


def test_login_identity_limit_spans_source_addresses(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(auth_settings, "auth_rate_limit", 2)
    payload = {"email": "target@cyberpass.example.com", "password": "Bbbbbbbbbbbb2"}
    with (
        TestClient(fastapi_app, client=("198.51.100.10", 50000)) as first_ip,
        TestClient(fastapi_app, client=("203.0.113.20", 50000)) as second_ip,
    ):
        assert first_ip.post("/api/v1/auth/login", json=payload).status_code == 401
        assert second_ip.post("/api/v1/auth/login", json=payload).status_code == 401
        limited = second_ip.post("/api/v1/auth/login", json=payload)
    assert limited.status_code == 429
    assert int(limited.headers["retry-after"]) >= 1
